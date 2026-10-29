"""excel_export.py -- Generacion del Excel consolidado por subcuenta a
partir de las filas parseadas. Usa openpyxl para estilos. Separado de
historial_facturas_ute.py sin cambios de logica.
"""
import os
from datetime import datetime

import pandas as pd

from parsing_facturas import _to_excel_number


# =========================
# Guardar Excel consolidado
# =========================
def guardar_excel_consolidado(filas: list[dict], cuenta: str, output_dir: str, log):
    if not filas:
        log(f"  [{cuenta}] Sin datos para exportar.")
        return None
 
    df = pd.DataFrame(filas)
 
    COLS_TEXTO = {"nro_cuenta", "nro_medidor", "acuerdo_servicio", "nro_factura"}
    for col in COLS_TEXTO:
        if col in df.columns:
            df[col] = df[col].fillna("").astype(str).str.strip()
            df[col] = df[col].str.replace(r"\.0$", "", regex=True)

    COLS_NO_NUMERICAS = COLS_TEXTO | {
        "archivo",
        "tipo_factura",
        "fecha_emision",
        "prox_vencimiento",
        "tarifa_aplicada",
        "potencia_contratada_punta_llano_kw",
        "tension",
        "fases",
        "direccion_servicio",
        "nombre_cliente",
        "direccion_cliente",
        "localidad_cliente",
        "departamento_cliente",
        "periodo_consumo",
        "zona_electrica",
    }
    for col in df.columns:
        if col in COLS_NO_NUMERICAS or col.endswith("_tipo_lec"):
            continue
        df[col] = df[col].apply(_to_excel_number)
 
    cols_fijas = [
        "archivo", "tipo_factura", "nro_cuenta", "nro_factura", "fecha_emision", "prox_vencimiento",
        "acuerdo_servicio", "nro_medidor", "tarifa_aplicada",
        "nombre_cliente", "direccion_cliente", "localidad_cliente", "departamento_cliente",
        "potencia_contratada_punta_llano_kw", "potencia_contratada_valle_kw", "potencia_contratada_kw",
        "consumo_activo_kwh", "consumo_reactivo_kvarh",
        "tension", "fases", "direccion_servicio", "periodo_consumo", "zona_electrica",
        "total_detalle_facturacion",
    ]
    cols_lecturas = sorted(c for c in df.columns if c not in cols_fijas)
    df = df[[c for c in cols_fijas if c in df.columns] + cols_lecturas]
 
    # Sufijo TH/GS segun el tipo de factura predominante del suministro (2
    # hojas o 1 hoja). Todas las facturas de una misma cuenta deberian ser del
    # mismo tipo; se usa el mas frecuente por si alguna quedo mal detectada.
    tipos_factura = [str(f.get("tipo_factura") or "").strip().upper() for f in filas]
    tipos_factura = [t for t in tipos_factura if t in ("TH", "GS")]
    tipo_factura_cuenta = max(set(tipos_factura), key=tipos_factura.count) if tipos_factura else "TH"

    cuenta_nombre = f"{_safe_path_fragment(str(cuenta))} {tipo_factura_cuenta}"
    excel_out = os.path.join(output_dir, f"{cuenta_nombre}.xlsx")
    os.makedirs(os.path.dirname(excel_out), exist_ok=True)

    # Si ya existe y está bloqueado, eliminarlo antes de escribir
    if os.path.exists(excel_out):
        try:
            os.remove(excel_out)
        except PermissionError:
            # Archivo abierto en Windows: renombrar el viejo y escribir el nuevo
            ts = datetime.now().strftime("%Y%m%d_%H%M%S")
            excel_viejo = os.path.join(output_dir, f"{cuenta_nombre}_OLD_{ts}.xlsx")
            try:
                os.rename(excel_out, excel_viejo)
                log(f"  [{cuenta}] [INFO] Archivo anterior renombrado a: {os.path.basename(excel_viejo)}")
            except Exception:
                pass  # Si tampoco se puede renombrar, intentar escribir igual
 
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.styles.numbers import FORMAT_TEXT
 
    with pd.ExcelWriter(excel_out, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Historial", index=False)
        ws = writer.sheets["Historial"]
 
        col_idx = {cell.value: cell.column for cell in ws[1]}
        for col_name in COLS_TEXTO:
            if col_name in col_idx:
                c = col_idx[col_name]
                for row in ws.iter_rows(min_row=2, min_col=c, max_col=c):
                    for cell in row:
                        cell.number_format = FORMAT_TEXT
                        if cell.value is not None:
                            cell.value = str(cell.value).replace(".0", "")
 
        hf = PatternFill("solid", start_color="1F4E79", end_color="1F4E79")
        hfont = Font(bold=True, color="FFFFFF", name="Arial", size=10)
        dfont = Font(name="Arial", size=10)
 
        for cell in ws[1]:
            cell.font = hfont
            cell.fill = hf
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
 
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = dfont
                if isinstance(cell.value, int):
                    cell.alignment = Alignment(horizontal="right")
                    cell.number_format = "0"
                elif isinstance(cell.value, float):
                    cell.alignment = Alignment(horizontal="right")
                    cell.number_format = "0.00"
                else:
                    cell.alignment = Alignment(horizontal="left")
 
        for col in ws.columns:
            max_len = max((len(str(cell.value or "")) for cell in col), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)
 
        ws.row_dimensions[1].height = 35
 
    log(f"  [{cuenta}] Excel guardado: {excel_out}")
    return excel_out
 
 
