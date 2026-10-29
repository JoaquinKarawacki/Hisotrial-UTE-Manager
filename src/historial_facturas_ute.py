from __future__ import annotations
 
import os
import re
import time
from datetime import datetime
from urllib.parse import urljoin
 
import pandas as pd
import pdfplumber
import requests
from playwright.sync_api import sync_playwright, TimeoutError as PlaywrightTimeoutError
from parsing_facturas import (
    es_pdf_real,
    contar_paginas_pdf,
    detectar_tipo_factura,
    parse_factura_historial,
    _to_excel_number,
)
 
# =========================
# CONFIG
# =========================
UTE_LOGIN = "https://identityserver.ute.com.uy/Account/Login"
SERVICIOS_EN_LINEA = "https://www.ute.com.uy/clientes/tramites-y-servicios/servicios-en-linea"
 
MAX_FACTURAS_VALIDAS = 24
 
 
# =========================
# Descarga PDF con cookies
# =========================
def cookies_to_requests(context) -> requests.Session:
    s = requests.Session()
    for c in context.cookies():
        s.cookies.set(c["name"], c["value"], domain=c.get("domain"), path=c.get("path"))
    s.headers.update({
        "User-Agent": "Mozilla/5.0",
        "Accept": "application/pdf,application/octet-stream;q=0.9,*/*;q=0.8",
        "Accept-Language": "es-ES,es;q=0.9,en;q=0.8",
        "Connection": "keep-alive",
    })
    return s
 
 
def descargar_pdf_con_cookies(context, url: str, out_pdf: str) -> str | None:
    s = cookies_to_requests(context)
    try:
        r = s.get(url, allow_redirects=True, timeout=60)
        if r.status_code != 200:
            return None
        data = r.content
        if data[:5] != b"%PDF-":
            return None
        os.makedirs(os.path.dirname(out_pdf), exist_ok=True)
        with open(out_pdf, "wb") as f:
            f.write(data)
        return out_pdf
    except Exception:
        return None
 
 
def extraer_posibles_urls_pdf(html: str, base_url: str) -> list[str]:
    candidates = set()
    for m in re.findall(r"""href=["']([^"']+\.pdf[^"']*)["']""", html, flags=re.I):
        candidates.add(urljoin(base_url, m))
    for m in re.findall(r"""src=["']([^"']+\.pdf[^"']*)["']""", html, flags=re.I):
        candidates.add(urljoin(base_url, m))
    for m in re.findall(r"""(https?://[^\s"']+\.pdf[^\s"']*)""", html, flags=re.I):
        candidates.add(m)
    return list(candidates)
 
 
# =========================
# Navegacion UTE — helpers
# =========================
def esperar_frame_autoservicio(page):
    patrones_url = [
        re.compile(r"autoservicio\.ute\.com\.uy", re.I),
        re.compile(r"selfservice", re.I),
    ]
    for _ in range(160):
        for fr in page.frames:
            fr_url = (fr.url or "").strip()
            if fr_url and any(p.search(fr_url) for p in patrones_url):
                return fr
        for sel in [
            "iframe[src*='autoservicio']",
            "iframe[src*='SelfService']",
            "iframe[name*='autoservicio' i]",
            "iframe[id*='autoservicio' i]",
            "iframe[src*='ute.com.uy']",
        ]:
            loc = page.locator(sel)
            if loc.count() > 0:
                try:
                    fr = loc.first.content_frame()
                    if fr:
                        return fr
                except Exception:
                    pass
        page.wait_for_timeout(250)
    return None
 
 
def _resumen_frames_debug(page, max_items: int = 8) -> str:
    frames = []
    for fr in page.frames:
        try:
            u = (fr.url or "").strip()
        except Exception:
            u = ""
        if u:
            frames.append(u)
    vistos = list(dict.fromkeys(frames))
    if len(vistos) > max_items:
        return " | ".join(vistos[:max_items]) + f" | ... (+{len(vistos) - max_items})"
    return " | ".join(vistos) if vistos else "(sin frames con URL)"
 
 
def _es_error_frame_detached(exc: Exception) -> bool:
    msg = str(exc).lower()
    return (
        "frame was detached" in msg
        or "execution context was destroyed" in msg
        or "target page, context or browser has been closed" in msg
    )
 
 
def _reenganchar_frame(page, frame):
    try:
        if frame and not frame.is_detached():
            return frame
    except Exception:
        pass
    return esperar_frame_autoservicio(page)


def _esperar_tabla_libre(frame, timeout_ms: int = 12000):
    """Espera que el overlay de loading de jTable desaparezca antes de hacer click."""
    try:
        frame.locator("div.jtable-busy-panel-background").wait_for(state="hidden", timeout=timeout_ms)
    except Exception:
        pass
 
 
def recolectar_todas_las_cuentas(frame, page) -> list[str]:
    cuentas = []
    pagina = 1
    reintentos_detach = 0
    while True:
        try:
            frame = _reenganchar_frame(page, frame)
            if not frame:
                break
            frame.locator("tr.jtable-data-row").first.wait_for(state="visible", timeout=15000)
            filas = frame.locator("tr.jtable-data-row")
            for i in range(filas.count()):
                numero = filas.nth(i).locator("td").nth(0).inner_text().strip()
                if numero and numero not in cuentas:
                    cuentas.append(numero)
            btn_next = frame.locator("span.jtable-page-number-next")
            if btn_next.count() == 0:
                break
            clase = btn_next.first.get_attribute("class") or ""
            if "disabled" in clase:
                break
            _esperar_tabla_libre(frame)
            btn_next.first.click()
            page.wait_for_timeout(1500)
            frame.locator("tr.jtable-data-row").first.wait_for(state="visible", timeout=10000)
            pagina += 1
            reintentos_detach = 0
        except Exception as e:
            if _es_error_frame_detached(e) and reintentos_detach < 4:
                reintentos_detach += 1
                page.wait_for_timeout(900)
                frame = esperar_frame_autoservicio(page)
                continue
            raise
    return cuentas
 
 
def _seleccionar_cuenta_en_tabla(frame, page, cuenta_actual: str, max_paginas: int = 80) -> bool:
    paginas_recorridas = 0
    reintentos_detach = 0
    while paginas_recorridas < max_paginas:
        try:
            frame = _reenganchar_frame(page, frame)
            if not frame:
                return False
            frame.locator("tr.jtable-data-row").first.wait_for(state="visible", timeout=12000)
        except PlaywrightTimeoutError:
            return False
        except Exception as e:
            if _es_error_frame_detached(e) and reintentos_detach < 4:
                reintentos_detach += 1
                page.wait_for_timeout(900)
                frame = esperar_frame_autoservicio(page)
                continue
            return False
 
        filas = frame.locator("tr.jtable-data-row")
        try:
            total_filas = filas.count()
        except Exception:
            return False
 
        for j in range(total_filas):
            try:
                txt = filas.nth(j).locator("td").nth(0).inner_text(timeout=3000).strip()
            except Exception:
                continue
            if txt == cuenta_actual:
                try:
                    filas.nth(j).click(timeout=5000)
                    return True
                except Exception:
                    return False
 
        btn_next = frame.locator("span.jtable-page-number-next")
        if btn_next.count() == 0:
            return False
        clase = btn_next.first.get_attribute("class") or ""
        if "disabled" in clase:
            return False
        _esperar_tabla_libre(frame)
        btn_next.first.click()
        page.wait_for_timeout(900)
        paginas_recorridas += 1
        reintentos_detach = 0
 
    return False
 
 
def _esperar_contexto_subcuenta(page, frame, timeout_ms: int = 30000) -> bool:
    intentos = max(1, timeout_ms // 500)
    for _ in range(intentos):
        try:
            frame = _reenganchar_frame(page, frame)
            if not frame:
                page.wait_for_timeout(500)
                continue
            if frame.locator("text=/Informaci[oó]n\\s+de\\s+la\\s+cuenta/i").count() > 0:
                return True
            if frame.locator("text=/N[uú]mero\\s+de\\s+cuenta/i").count() > 0:
                return True
            if frame.locator("text=/Mis\\s+Facturas/i").count() > 0:
                return True
        except Exception as e:
            if not _es_error_frame_detached(e):
                pass
        page.wait_for_timeout(500)
    return False
 
 
def _goto_estable(page, url: str, wait_until: str = "domcontentloaded", timeout: int = 60000, reintentos: int = 5) -> bool:
    for _ in range(reintentos):
        try:
            page.goto(url, wait_until=wait_until, timeout=timeout)
            return True
        except Exception as e:
            msg = str(e).lower()
            if "interrupted by another navigation" in msg or "net::err_abort" in msg:
                try:
                    page.wait_for_load_state("domcontentloaded", timeout=10000)
                except Exception:
                    pass
                page.wait_for_timeout(1200)
                continue
            raise
    return False
 
 
def _detectar_modo_multicuentas(page, frame, timeout_ms: int = 25000) -> bool:
    intentos = max(1, timeout_ms // 500)
    for _ in range(intentos):
        try:
            frame = _reenganchar_frame(page, frame)
            if not frame:
                page.wait_for_timeout(500)
                continue
            filas = frame.locator("tr.jtable-data-row")
            if filas.count() > 0:
                return True
            hay_numero_cuenta = frame.locator("text=/N[uú]mero\\s+de\\s+cuenta/i").count() > 0
            hay_mis_facturas = frame.locator("text=/Mis\\s+Facturas/i").count() > 0
            if hay_numero_cuenta and hay_mis_facturas:
                return False
        except Exception as e:
            if not _es_error_frame_detached(e):
                pass
        page.wait_for_timeout(500)
    return True
 
 
# =========================
# Mis Facturas — Historial completo
# =========================

def abrir_mis_facturas_historial(page, frame, log):
    log("  Haciendo clic en 'Mis Facturas'...")
    frame.get_by_text("Mis Facturas").click()
    page.wait_for_timeout(1000)
 
    activado = False
    try:
        sel = frame.locator("select#getAll, select[name='getAll']")
        sel1 = frame.locator("Select#tipoDoc, select[name='tipoDoc']")
        sel.wait_for(state="visible", timeout=8000)
        sel1.wait_for(state="visible", timeout=8000)
        sel.select_option(value="true")
        sel1.select_option(label="Factura")
        page.wait_for_timeout(600)
        log("Solo facturas y Historial completo activado.")
        activado = True
    except PlaywrightTimeoutError:
        pass
    except Exception as e:
        log(f"  [WARN] select#getAll: {e}")
 
    if not activado:
        try:
            sel = frame.locator("select").filter(has_text=re.compile(r"[Hh]istorial"))
            sel1 = frame.locator("select").filter(has=frame.locator("option", has_text=re.compile(r"Factura", re.I)))
            if sel.count() > 0:
                sel.first.select_option(value="true")
                sel1.first.select_option(label="Factura")
                page.wait_for_timeout(600)
                log("Solo factursa y Historial completo activado (fallback).")
                activado = True
        except Exception:
            pass
 
    if not activado:
        log("  [WARN] No se pudo activar el filtrado de solo facturas y historial completo.")
 
    try:
        btn_buscar = frame.locator("#btnBuscar, button#btnBuscar")
        btn_buscar.first.wait_for(state="visible", timeout=8000)
        btn_buscar.first.click()
        page.wait_for_timeout(1500)
    except Exception:
        try:
            frame.get_by_role("button", name=re.compile("Buscar", re.I)).click()
            page.wait_for_timeout(1500)
        except Exception:
            pass
 
 
def _scroll_tabla_historial(frame) -> bool:
    candidates = [
        ".jtable-main-container",
        ".jtable-main-container > div",
        ".table-responsive",
        ".jtable",
        "#tblHistorialFacturas",
    ]
    for sel in candidates:
        try:
            box = frame.locator(sel)
            if box.count() == 0:
                continue
            moved = box.first.evaluate(
                """
                (el) => {
                    const maxTop = el.scrollHeight - el.clientHeight;
                    if (maxTop <= 0) return false;
                    const prev = el.scrollTop;
                    const step = Math.max(220, Math.floor(el.clientHeight * 0.8));
                    el.scrollTop = Math.min(maxTop, prev + step);
                    return el.scrollTop > prev;
                }
                """
            )
            if moved:
                return True
        except Exception:
            continue
    try:
        moved = frame.evaluate(
            """
            () => {
                const doc = document.scrollingElement || document.documentElement || document.body;
                if (!doc) return false;
                const maxTop = doc.scrollHeight - doc.clientHeight;
                if (maxTop <= 0) return false;
                const prev = doc.scrollTop;
                const step = Math.max(220, Math.floor(doc.clientHeight * 0.8));
                doc.scrollTop = Math.min(maxTop, prev + step);
                return doc.scrollTop > prev;
            }
            """
        )
        return bool(moved)
    except Exception:
        return False
 
 
def obtener_links_facturas_historial(frame, min_links: int = MAX_FACTURAS_VALIDAS) -> list[str]:
    try:
        frame.locator("a[href*='cmviewbill']").first.wait_for(state="attached", timeout=30000)
    except PlaywrightTimeoutError:
        return []
 
    hrefs = []
    sin_cambios = 0
    for _ in range(35):
        links = frame.locator("a[href*='cmviewbill']").filter(has_text=re.compile(r"\d"))
        before = len(hrefs)
        for i in range(links.count()):
            href = links.nth(i).get_attribute("href")
            if href:
                full = urljoin(frame.url, href)
                if full not in hrefs:
                    hrefs.append(full)
        if len(hrefs) >= min_links:
            break
        if len(hrefs) == before:
            sin_cambios += 1
        else:
            sin_cambios = 0
        if sin_cambios >= 3:
            break
        if not _scroll_tabla_historial(frame):
            break
        frame.wait_for_timeout(280)
 
    return hrefs
 
 
def hay_mas_paginas_historial(frame) -> bool:
    try:
        candidatos = [
            "span.jtable-page-number-next",
            "a.jtable-page-number-next",
            "li.next a",
            "a[aria-label='Next']",
            "a[title='Next']",
        ]
        for sel in candidatos:
            btn = frame.locator(sel)
            if btn.count() == 0:
                continue
            clase = (btn.first.get_attribute("class") or "").lower()
            aria_disabled = (btn.first.get_attribute("aria-disabled") or "").lower()
            parent_class = ""
            try:
                parent_class = (btn.first.evaluate("el => (el.parentElement && el.parentElement.className) || ''") or "").lower()
            except Exception:
                pass
            if "disabled" in clase or "disabled" in parent_class or aria_disabled == "true":
                continue
            return True
        return False
    except Exception:
        return False
 
 
def ir_a_siguiente_pagina_historial(frame, page, primer_href_anterior: str = "") -> bool:
    try:
        pagina_actual = None
        try:
            act = frame.locator(".jtable-page-number-active, .pagination .active")
            if act.count() > 0:
                pagina_actual = (act.first.inner_text() or "").strip()
        except Exception:
            pagina_actual = None
 
        if not primer_href_anterior:
            try:
                first_link = frame.locator("a[href*='cmviewbill']")
                if first_link.count() > 0:
                    primer_href_anterior = first_link.first.get_attribute("href") or ""
            except Exception:
                primer_href_anterior = ""
 
        clickeado = False
        for sel in [
            "span.jtable-page-number-next",
            "a.jtable-page-number-next",
            "li.next a",
            "a[aria-label='Next']",
            "a[title='Next']",
        ]:
            btn = frame.locator(sel)
            if btn.count() == 0:
                continue
            clase = (btn.first.get_attribute("class") or "").lower()
            aria_disabled = (btn.first.get_attribute("aria-disabled") or "").lower()
            parent_class = ""
            try:
                parent_class = (btn.first.evaluate("el => (el.parentElement && el.parentElement.className) || ''") or "").lower()
            except Exception:
                pass
            if "disabled" in clase or "disabled" in parent_class or aria_disabled == "true":
                continue
            _esperar_tabla_libre(frame)
            btn.first.click(timeout=6000)
            clickeado = True
            break
 
        if not clickeado:
            return False
 
        for _ in range(20):
            page.wait_for_timeout(300)
            frame = _reenganchar_frame(page, frame)
            if not frame:
                continue
            cambio_pagina = False
            try:
                if pagina_actual is not None:
                    act = frame.locator(".jtable-page-number-active, .pagination .active")
                    if act.count() > 0:
                        pagina_nueva = (act.first.inner_text() or "").strip()
                        if pagina_nueva and pagina_nueva != pagina_actual:
                            cambio_pagina = True
            except Exception:
                pass
            cambio_href = False
            try:
                first_link = frame.locator("a[href*='cmviewbill']")
                if first_link.count() > 0:
                    href_nuevo = first_link.first.get_attribute("href") or ""
                    if href_nuevo and href_nuevo != primer_href_anterior:
                        cambio_href = True
            except Exception:
                pass
            if cambio_pagina or cambio_href:
                return True
 
        return False
    except Exception:
        return False
 
 
# =========================
# Procesamiento por subcuenta
# =========================
def descargar_factura_por_url(context, page, bill_url: str, out_pdf: str, log) -> str | None:
    os.makedirs(os.path.dirname(out_pdf), exist_ok=True)
    pdf_path = descargar_pdf_con_cookies(context, bill_url, out_pdf)
    if pdf_path:
        return pdf_path
    try:
        page.goto(bill_url, wait_until="domcontentloaded", timeout=60000)
        page.wait_for_timeout(1200)
        html = page.content()
        for u in extraer_posibles_urls_pdf(html, page.url):
            pdf_path = descargar_pdf_con_cookies(context, u, out_pdf)
            if pdf_path:
                return pdf_path
    except Exception as e:
        log(f"    [WARN] {e}")
    return None


def _safe_path_fragment(value: str) -> str:
    """Normaliza texto para usarlo en nombres de archivo/carpeta en Windows."""
    raw = str(value or "").strip()
    cleaned = re.sub(r'[\\/:*?"<>|]+', "_", raw)
    return cleaned or "sin_nombre"


def _renombrar_factura(pdf_path: str, cuenta_actual, datos: dict, n: int) -> str:
    """Renombra 'factura_NNN.pdf' a '{cuenta} {GS|TH} {NNN}_{nro_factura}.pdf'.

    El correlativo n (orden de descarga, que respeta el orden del historial
    de UTE: la mas reciente primero) siempre va en el nombre para que al
    ordenar la carpeta alfabeticamente la 001 sea la factura mas reciente y
    la ultima (024) la mas vieja, igual que antes de nombrar los PDF segun
    el nro_factura. Si no se pudo parsear nro_factura, se omite ese sufijo.
    """
    tipo_factura = datos.get("tipo_factura") or "TH"
    nro_factura_raw = (datos.get("nro_factura") or "").replace(" ", "")
    ident = _safe_path_fragment(nro_factura_raw) if nro_factura_raw else ""
    out_dir = os.path.dirname(pdf_path)
    base = f"{_safe_path_fragment(str(cuenta_actual))} {tipo_factura} {n:03d}"
    nuevo_nombre = f"{base}_{ident}.pdf" if ident else f"{base}.pdf"
    nuevo_path = os.path.join(out_dir, nuevo_nombre)

    contador = 1
    while os.path.exists(nuevo_path) and os.path.abspath(nuevo_path) != os.path.abspath(pdf_path):
        nuevo_path = os.path.join(out_dir, f"{nuevo_nombre[:-4]}_dup{contador}.pdf")
        contador += 1

    try:
        os.replace(pdf_path, nuevo_path)
        return nuevo_path
    except OSError:
        return pdf_path
 
 
def procesar_subcuenta_historial(context, page, frame, cuenta_actual, output_dir, log) -> list[dict]:
    abrir_mis_facturas_historial(page, frame, log)
 
    cuenta_nombre = _safe_path_fragment(str(cuenta_actual))
    out_dir = os.path.join(output_dir, f"historial_{cuenta_nombre}")
    os.makedirs(out_dir, exist_ok=True)
 
    validas: list[dict] = []
    filas_excel: list[dict] = []
    descartadas = 0
    pagina = 1
    inicio = time.monotonic()
    max_segundos = 12 * 60
    max_intentos_descarga = 180
    intentos_descarga = 0
    urls_vistas: set[str] = set()
    paginas_sin_nuevos = 0
    firma_pagina_anterior: tuple[str, ...] = ()
    repeticiones_misma_pagina = 0
 
    while len(validas) < MAX_FACTURAS_VALIDAS:
        if (time.monotonic() - inicio) > max_segundos:
            log(f"  [{cuenta_actual}] Corte por tiempo. Se exporta parcial.")
            break
        if intentos_descarga >= max_intentos_descarga:
            log(f"  [{cuenta_actual}] Corte por intentos. Se exporta parcial.")
            break
 
        frame = _reenganchar_frame(page, frame)
        if not frame:
            break
 
        links = obtener_links_facturas_historial(frame)
        if not links:
            log(f"  [{cuenta_actual}] Sin links en pagina {pagina}.")
            break

        firma_pagina_actual = tuple(links[:5])
        if firma_pagina_actual and firma_pagina_actual == firma_pagina_anterior:
            repeticiones_misma_pagina += 1
        else:
            repeticiones_misma_pagina = 0
            firma_pagina_anterior = firma_pagina_actual

        if repeticiones_misma_pagina >= 1:
            log(f"  [{cuenta_actual}] Fin del historial detectado: la pagina no cambia al avanzar.")
            break

        links_nuevos = [u for u in links if u not in urls_vistas]
        if not links_nuevos:
            log(f"  [{cuenta_actual}] Pag {pagina}: sin links nuevos.")
            paginas_sin_nuevos += 1
        else:
            log(f"  [{cuenta_actual}] Pag {pagina}: {len(links_nuevos)} facturas nuevas.")
            paginas_sin_nuevos = 0

        for bill_url in links_nuevos:
            if len(validas) >= MAX_FACTURAS_VALIDAS:
                break
            if (time.monotonic() - inicio) > max_segundos:
                break
            if intentos_descarga >= max_intentos_descarga:
                break
 
            urls_vistas.add(bill_url)
            intentos_descarga += 1
 
            n = len(validas) + descartadas + 1
            out_pdf = os.path.join(out_dir, f"factura_{n:03d}.pdf")
 
            pdf_path = descargar_factura_por_url(context, page, bill_url, out_pdf, log)
            if not pdf_path or not es_pdf_real(pdf_path):
                descartadas += 1
                continue
 
            npags = contar_paginas_pdf(pdf_path)
            if npags <= 0:
                log(f"    [{cuenta_actual}] PDF invalido — descartada.")
                descartadas += 1
                continue

            try:
                datos = parse_factura_historial(pdf_path)
                parse_ok = True
            except Exception as e:
                log(f"    [{cuenta_actual}] Error parseando ({npags} pag(s)): {e}")
                datos = {"tipo_factura": detectar_tipo_factura(pdf_path)}
                parse_ok = False

            # El renombrado no depende de que el parseo haya sido exitoso: aun
            # si parse_factura_historial fallo, se dispone del tipo (TH/GS) y
            # el archivo no debe quedar con el nombre temporal "factura_NNN.pdf".
            datos["nro_cuenta"] = datos.get("nro_cuenta") or cuenta_actual
            datos["npags_pdf"] = npags
            pdf_path = _renombrar_factura(pdf_path, cuenta_actual, datos, n)
            datos["archivo"] = os.path.basename(pdf_path)

            if parse_ok:
                validas.append(datos)
                filas_excel.append(datos)
                log(f"    [{cuenta_actual}] {len(validas)}/{MAX_FACTURAS_VALIDAS} validas ({npags} pag(s)).")
            else:
                descartadas += 1

        if len(validas) >= MAX_FACTURAS_VALIDAS:
            break
        if (time.monotonic() - inicio) > max_segundos:
            break
        if intentos_descarga >= max_intentos_descarga:
            break

        if paginas_sin_nuevos >= 2:
            log(f"  [{cuenta_actual}] Fin del historial detectado: no aparecen facturas nuevas.")
            break

        primer_href_actual = links[0] if links else ""
        if not hay_mas_paginas_historial(frame):
            log(f"  [{cuenta_actual}] Fin del historial detectado: no hay mas paginas.")
            break

        if not ir_a_siguiente_pagina_historial(frame, page, primer_href_actual):
            log(f"  [{cuenta_actual}] No se pudo avanzar a la siguiente pagina del historial.")
            break

        pagina += 1
 
    log(
        f"  [{cuenta_actual}] Validas: {len(validas)} | Descartadas: {descartadas} | "
        f"Intentos: {intentos_descarga}"
    )
    return filas_excel
 
 
# =========================
# Guardar Excel consolidado
# =========================
def guardar_excel_consolidado(filas: list[dict], cuenta: str, output_dir: str, log):
    if not filas:
        log(f"  [{cuenta}] Sin datos para exportar.")
        return None
 
    df = pd.DataFrame(filas)
 
    COLS_TEXTO = {"nro_cuenta", "nro_medidor", "acuerdo_servicio", "nro_factura"}
    for col in COLS_TEXTO:
        if col in df.columns:
            df[col] = df[col].fillna("").astype(str).str.strip()
            df[col] = df[col].str.replace(r"\.0$", "", regex=True)

    COLS_NO_NUMERICAS = COLS_TEXTO | {
        "archivo",
        "tipo_factura",
        "fecha_emision",
        "prox_vencimiento",
        "tarifa_aplicada",
        "potencia_contratada_punta_llano_kw",
        "tension",
        "fases",
        "direccion_servicio",
        "nombre_cliente",
        "direccion_cliente",
        "localidad_cliente",
        "departamento_cliente",
        "periodo_consumo",
        "zona_electrica",
    }
    for col in df.columns:
        if col in COLS_NO_NUMERICAS or col.endswith("_tipo_lec"):
            continue
        df[col] = df[col].apply(_to_excel_number)
 
    cols_fijas = [
        "archivo", "tipo_factura", "nro_cuenta", "nro_factura", "fecha_emision", "prox_vencimiento",
        "acuerdo_servicio", "nro_medidor", "tarifa_aplicada",
        "nombre_cliente", "direccion_cliente", "localidad_cliente", "departamento_cliente",
        "potencia_contratada_punta_llano_kw", "potencia_contratada_valle_kw", "potencia_contratada_kw",
        "consumo_activo_kwh", "consumo_reactivo_kvarh",
        "tension", "fases", "direccion_servicio", "periodo_consumo", "zona_electrica",
        "total_detalle_facturacion",
    ]
    cols_lecturas = sorted(c for c in df.columns if c not in cols_fijas)
    df = df[[c for c in cols_fijas if c in df.columns] + cols_lecturas]
 
    # Sufijo TH/GS segun el tipo de factura predominante del suministro (2
    # hojas o 1 hoja). Todas las facturas de una misma cuenta deberian ser del
    # mismo tipo; se usa el mas frecuente por si alguna quedo mal detectada.
    tipos_factura = [str(f.get("tipo_factura") or "").strip().upper() for f in filas]
    tipos_factura = [t for t in tipos_factura if t in ("TH", "GS")]
    tipo_factura_cuenta = max(set(tipos_factura), key=tipos_factura.count) if tipos_factura else "TH"

    cuenta_nombre = f"{_safe_path_fragment(str(cuenta))} {tipo_factura_cuenta}"
    excel_out = os.path.join(output_dir, f"{cuenta_nombre}.xlsx")
    os.makedirs(os.path.dirname(excel_out), exist_ok=True)

    # Si ya existe y está bloqueado, eliminarlo antes de escribir
    if os.path.exists(excel_out):
        try:
            os.remove(excel_out)
        except PermissionError:
            # Archivo abierto en Windows: renombrar el viejo y escribir el nuevo
            ts = datetime.now().strftime("%Y%m%d_%H%M%S")
            excel_viejo = os.path.join(output_dir, f"{cuenta_nombre}_OLD_{ts}.xlsx")
            try:
                os.rename(excel_out, excel_viejo)
                log(f"  [{cuenta}] [INFO] Archivo anterior renombrado a: {os.path.basename(excel_viejo)}")
            except Exception:
                pass  # Si tampoco se puede renombrar, intentar escribir igual
 
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.styles.numbers import FORMAT_TEXT
 
    with pd.ExcelWriter(excel_out, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Historial", index=False)
        ws = writer.sheets["Historial"]
 
        col_idx = {cell.value: cell.column for cell in ws[1]}
        for col_name in COLS_TEXTO:
            if col_name in col_idx:
                c = col_idx[col_name]
                for row in ws.iter_rows(min_row=2, min_col=c, max_col=c):
                    for cell in row:
                        cell.number_format = FORMAT_TEXT
                        if cell.value is not None:
                            cell.value = str(cell.value).replace(".0", "")
 
        hf = PatternFill("solid", start_color="1F4E79", end_color="1F4E79")
        hfont = Font(bold=True, color="FFFFFF", name="Arial", size=10)
        dfont = Font(name="Arial", size=10)
 
        for cell in ws[1]:
            cell.font = hfont
            cell.fill = hf
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
 
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = dfont
                if isinstance(cell.value, int):
                    cell.alignment = Alignment(horizontal="right")
                    cell.number_format = "0"
                elif isinstance(cell.value, float):
                    cell.alignment = Alignment(horizontal="right")
                    cell.number_format = "0.00"
                else:
                    cell.alignment = Alignment(horizontal="left")
 
        for col in ws.columns:
            max_len = max((len(str(cell.value or "")) for cell in col), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)
 
        ws.row_dimensions[1].height = 35
 
    log(f"  [{cuenta}] Excel guardado: {excel_out}")
    return excel_out
 
 
# =========================
# FUNCION PRINCIPAL
# =========================
def run_historial(
    ute_user: str,
    ute_pass: str,
    cuentas_seleccionadas: list[str],
    output_dir: str,
    log_fn=None,
):
    def log(msg):
        if log_fn:
            log_fn(msg)
        else:
            print(msg)
 
    if not ute_user or not ute_pass:
        log("ERROR: credenciales vacias.")
        return
 
    os.makedirs(output_dir, exist_ok=True)
    log(f"Guardando en: {output_dir}")
 
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=False)
        context = browser.new_context(viewport={"width": 1600, "height": 900}, accept_downloads=True)
        context.set_default_timeout(15000)
        context.set_default_navigation_timeout(60000)
        page = context.new_page()
        page.set_default_timeout(15000)
        page.set_default_navigation_timeout(60000)
 
        try:
            log("Iniciando sesion en UTE...")
            if not _goto_estable(page, UTE_LOGIN, wait_until="domcontentloaded", timeout=60000):
                log("No se pudo abrir login de UTE.")
                return
            page.wait_for_selector("input#Username", timeout=60000)
            page.locator("input#Username").fill(ute_user)
            page.locator("input#Password").fill(ute_pass)
            btn = page.get_by_role("button", name=re.compile(r"^Ingresar$", re.I))
            btn.wait_for(state="visible", timeout=30000)
            # FIX: no_wait_after=True para evitar timeout por navegacion post-login
            btn.click(timeout=20000, no_wait_after=True)
            try:
                page.wait_for_load_state("networkidle", timeout=8000)
            except PlaywrightTimeoutError:
                pass
            page.wait_for_timeout(1200)
 
            if not _goto_estable(page, SERVICIOS_EN_LINEA, wait_until="domcontentloaded", timeout=60000):
                log("No se pudo abrir Servicios en linea.")
                return
            page.wait_for_timeout(3000)
 
            frame = esperar_frame_autoservicio(page)
            if not frame:
                log("No se encontro iframe de autoservicio.")
                log(f"Frames detectados: {_resumen_frames_debug(page)}")
                return
 
            multi = _detectar_modo_multicuentas(page, frame, timeout_ms=25000)
 
            if multi:
                log("Modo MULTI-CUENTAS.")
                todas = recolectar_todas_las_cuentas(frame, page)
                log(f"Subcuentas en portal: {len(todas)}")
                if not todas:
                    log("No se pudieron listar subcuentas. Probando CUENTA UNICA.")
                    multi = False
 
            if multi:
                a_procesar = (
                    [c for c in todas if c in cuentas_seleccionadas]
                    if cuentas_seleccionadas else todas
                )
                log(f"A procesar: {len(a_procesar)}")
 
                for cuenta_actual in a_procesar:
                    log(f"\n{'='*50}\n[{cuenta_actual}] Iniciando...")
                    datos: list[dict] = []
                    try:
                        if not _goto_estable(page, SERVICIOS_EN_LINEA, wait_until="domcontentloaded", timeout=60000):
                            log(f"[{cuenta_actual}] No se pudo abrir Servicios en linea.")
                            continue
                        page.wait_for_timeout(2500)
                        frame = esperar_frame_autoservicio(page)
                        if not frame:
                            log(f"[{cuenta_actual}] Sin iframe.")
                            continue
                        if not _seleccionar_cuenta_en_tabla(frame, page, cuenta_actual):
                            log(f"[{cuenta_actual}] No encontrada en tabla.")
                            continue
                        if not _esperar_contexto_subcuenta(page, frame, timeout_ms=30000):
                            log(f"[{cuenta_actual}] Timeout esperando carga.")
                            continue
                        page.wait_for_timeout(500)
                        frame = _reenganchar_frame(page, frame)
                        datos = procesar_subcuenta_historial(
                            context, page, frame, cuenta_actual, output_dir, log
                        )
                    except PlaywrightTimeoutError:
                        log(f"[{cuenta_actual}] Timeout.")
                    except Exception as e:
                        if _es_error_frame_detached(e):
                            log(f"[{cuenta_actual}] Frame recargado.")
                        else:
                            log(f"[{cuenta_actual}] Error: {e}")
                    finally:
                        try:
                            guardar_excel_consolidado(datos, cuenta_actual, output_dir, log)
                        except Exception as e:
                            log(f"[{cuenta_actual}] Error guardando Excel: {e}")
 
            if not multi:
                log("Modo CUENTA UNICA.")
                try:
                    cuenta_actual = (
                        frame.locator("td:has-text('Numero de cuenta') + td")
                        .first.inner_text().strip()
                    )
                except Exception:
                    cuenta_actual = ute_user
                if not _esperar_contexto_subcuenta(page, frame, timeout_ms=30000):
                    log("Timeout esperando carga de cuenta unica.")
                    return
                page.wait_for_timeout(500)
                datos: list[dict] = []
                try:
                    datos = procesar_subcuenta_historial(
                        context, page, frame, cuenta_actual, output_dir, log
                    )
                finally:
                    guardar_excel_consolidado(datos, cuenta_actual, output_dir, log)
 
            log("\nProceso finalizado.")
 
        except Exception as e:
            if _es_error_frame_detached(e):
                log("Iframe recargado al finalizar.")
            else:
                log(f"Error inesperado: {e}")
        finally:
            browser.close()
 
 
# =========================
# MAIN — consola
# =========================
def _debug_parsear_pdfs(paths: list[str]):
    """Corre parse_factura_historial sobre PDFs locales e imprime el resultado.

    Util para validar cambios de parsing (GS/TH) sin correr el scraping real.
    Uso: python historial_facturas_ute.py debug "archivo1.pdf" "archivo2.pdf"
    """
    for path in paths:
        print("=" * 70)
        print(path)
        if not os.path.exists(path):
            print("  [ERROR] archivo no encontrado")
            continue
        datos = parse_factura_historial(path)
        for k, v in datos.items():
            print(f"  {k}: {v!r}")


if __name__ == "__main__":
    import sys as _sys
    if len(_sys.argv) > 1 and _sys.argv[1] == "debug":
        _debug_parsear_pdfs(_sys.argv[2:])
    else:
        ute_user = input("Usuario UTE: ").strip()
        ute_pass = input("Contrasena UTE: ").strip()
        output_dir = os.path.join(os.getcwd(), ute_user)
        run_historial(ute_user, ute_pass, cuentas_seleccionadas=[], output_dir=output_dir)
